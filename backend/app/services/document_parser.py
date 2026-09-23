import re
from io import BytesIO
from pathlib import Path

from docx import Document as WordDocument
from fastapi import HTTPException

from app.services.pdf_parser import extract_pdf_pages, normalize_text

SUPPORTED_FILE_TYPES = {
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}

LANGUAGE_NAMES = {"en": "English", "ar": "Arabic", "sw": "Swahili"}

_SWAHILI_WORDS = {
    "ambayo", "au", "bila", "cha", "hii", "hivyo", "ili", "katika", "kwa", "lakini",
    "na", "ni", "pia", "sana", "wa", "wakati", "ya", "yake", "yenye", "za",
}


def detect_language(text: str) -> str:
    letters = [character for character in text if character.isalpha()]
    if letters and sum("\u0600" <= character <= "\u06ff" for character in letters) / len(letters) >= 0.2:
        return "ar"
    words = re.findall(r"[a-z]+", text.casefold())
    if words:
        matches = sum(word in _SWAHILI_WORDS for word in words)
        if matches >= 3 and matches / len(words) >= 0.08:
            return "sw"
    return "en"


def extract_docx_pages(content: bytes) -> list[str]:
    if not content.startswith(b"PK"):
        raise HTTPException(422, "The file does not have a valid DOCX signature")
    try:
        document = WordDocument(BytesIO(content))
        blocks = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        text = normalize_text("\n".join(blocks))
    except Exception as exc:
        raise HTTPException(422, "The Word document could not be read; it may be corrupt") from exc
    if not text:
        raise HTTPException(422, "No readable text was found in the Word document")
    return [text]


def extract_document_pages(content: bytes, filename: str) -> tuple[list[str], str]:
    extension = Path(filename).suffix.lower()
    if extension == ".pdf":
        pages = extract_pdf_pages(content)
    elif extension == ".docx":
        pages = extract_docx_pages(content)
    else:
        raise HTTPException(415, "Upload a PDF or DOCX file")
    language = detect_language("\n".join(pages)[:20_000])
    return pages, language