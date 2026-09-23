from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
from fastapi import HTTPException
from pypdf import PdfWriter

from app.core.config import Settings
from app.services import document_parser
from app.services import pdf_parser as parser


def pdf_bytes(pages=1, encrypted=False):
    writer = PdfWriter()
    for _ in range(pages):
        writer.add_blank_page(width=612, height=792)
    if encrypted:
        writer.encrypt("test-password")
    stream = BytesIO()
    writer.write(stream)
    return stream.getvalue()


@pytest.fixture
def settings(monkeypatch):
    settings = Settings(_env_file=None)
    monkeypatch.setattr(parser, "get_settings", lambda: settings)
    return settings


@pytest.mark.parametrize("content,match", [(b"not a pdf", "signature"), (b"%PDF-broken", "corrupt")])
def test_corrupt_pdf(content, match, settings):
    with pytest.raises(HTTPException, match=match) as error:
        parser.extract_pdf_pages(content)
    assert error.value.status_code == 422


def test_encrypted_empty_and_page_limits(settings):
    for content in [pdf_bytes(encrypted=True), pdf_bytes(0)]:
        with pytest.raises(HTTPException) as error:
            parser.extract_pdf_pages(content)
        assert error.value.status_code == 422
    settings.pdf_max_pages = 1
    with pytest.raises(HTTPException, match="1 to 1"):
        parser.extract_pdf_pages(pdf_bytes(2))


def test_ocr_only_sparse_pages_preserves_page_alignment(settings, monkeypatch):
    reader = SimpleNamespace(is_encrypted=False, pages=[
        SimpleNamespace(extract_text=lambda: "Rich extracted text " * 10),
        SimpleNamespace(extract_text=lambda: ""),
        SimpleNamespace(extract_text=lambda: "Another rich page " * 10)])
    monkeypatch.setattr(parser, "PdfReader", lambda _: reader)
    monkeypatch.setattr("pdfplumber.open", Mock(side_effect=ValueError("Layout failure")))
    ocr = Mock(return_value="Recovered scan text")
    monkeypatch.setattr(parser, "_ocr_page", ocr)
    texts = parser.extract_pdf_pages(b"%PDF-fixture")
    assert len(texts) == 3
    assert texts[1] == "Recovered scan text"
    assert texts[0].startswith("Rich") and texts[2].startswith("Another")
    ocr.assert_called_once_with(b"%PDF-fixture", 1)


def test_ocr_page_budget_and_unreadable_scan(settings, monkeypatch):
    ocr = Mock(return_value="")
    monkeypatch.setattr(parser, "_ocr_page", ocr)
    settings.pdf_max_ocr_pages = 1
    with pytest.raises(HTTPException, match="split"):
        parser.extract_pdf_pages(pdf_bytes(2))
    ocr.assert_not_called()
    with pytest.raises(HTTPException, match="No readable text"):
        parser.extract_pdf_pages(pdf_bytes())


def test_normalize_tables_and_two_column_layout():
    assert parser.normalize_text(" heading \nheading\n row   text ") == "heading\nrow text"
    page = Mock(width=600, height=800)
    page.extract_tables.return_value = [[[" Name ", "Value"], [" Name ", "Value"], [None, " 10 "]]]
    assert parser._table_text(page) == "Name | Value\n| 10"
    page.extract_words.return_value = [{"x0": 10, "x1": 200}] * 8 + [{"x0": 400, "x1": 550}] * 8
    page.crop.side_effect = [SimpleNamespace(extract_text=lambda: "Left column"), SimpleNamespace(extract_text=lambda: "Right column")]
    assert parser._layout_text(page) == "Left column\nRight column"
    page.extract_text.assert_not_called()


def test_layout_recovers_page_and_deduplicates_table_rows(settings, monkeypatch):
    reader = SimpleNamespace(is_encrypted=False, pages=[SimpleNamespace(extract_text=lambda: "")])
    monkeypatch.setattr(parser, "PdfReader", lambda _: reader)
    page = Mock()
    page.extract_tables.return_value = [[["Carrier", "Energy required"], ["Carrier", "Energy required"]]]
    monkeypatch.setattr(parser, "_layout_text", lambda _: "Carrier Energy required\n" + "Active transport moves against gradients.")
    document = Mock()
    document.__enter__ = Mock(return_value=SimpleNamespace(pages=[page]))
    document.__exit__ = Mock(return_value=False)
    monkeypatch.setattr("pdfplumber.open", lambda _: document)
    ocr = Mock()
    monkeypatch.setattr(parser, "_ocr_page", ocr)
    text = parser.extract_pdf_pages(b"%PDF-fixture")[0]
    assert text.count("Carrier") == 1
    ocr.assert_not_called()


@pytest.mark.parametrize("fail", [False, True])
def test_ocr_render_bounds_and_cleanup(settings, monkeypatch, fail):
    import pypdfium2
    image, bitmap, page, document = Mock(), Mock(), Mock(), Mock()
    document.__getitem__ = Mock(return_value=page)
    page.get_size.return_value = (20000, 10000)
    page.render.return_value = bitmap
    bitmap.to_pil.return_value = image
    monkeypatch.setattr(pypdfium2, "PdfDocument", lambda _: document)
    engine = Mock(side_effect=RuntimeError("bad inference")) if fail else Mock(return_value=([[None, " OCR text ", 1]], None))
    monkeypatch.setattr(parser, "_ocr_engine", lambda: engine)
    if fail:
        with pytest.raises(HTTPException, match="page 1"):
            parser._ocr_page(b"pdf", 0)
    else:
        assert parser._ocr_page(b"pdf", 0) == "OCR text"
    scale = page.render.call_args.kwargs["scale"]
    assert 20000 * scale <= settings.pdf_ocr_max_dimension
    assert 20000 * 10000 * scale * scale <= settings.pdf_ocr_max_pixels
    for resource in [image, bitmap, page, document]:
        resource.close.assert_called_once()


def test_ocr_engine_is_cached(monkeypatch):
    import rapidocr_onnxruntime
    factory = Mock(return_value=object())
    parser._ocr_engine.cache_clear()
    monkeypatch.setattr(rapidocr_onnxruntime, "RapidOCR", factory)
    try:
        assert parser._ocr_engine() is parser._ocr_engine()
        factory.assert_called_once()
    finally:
        parser._ocr_engine.cache_clear()


def test_real_text_pdf_without_ocr(settings, monkeypatch):
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                             NameObject("/Subtype"): NameObject("/Type1"),
                             NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 40 700 Td (Diffusion moves particles down a concentration gradient.) Tj ET")
    page[NameObject("/Contents")] = stream
    output = BytesIO()
    writer.write(output)
    ocr = Mock(side_effect=AssertionError("Text PDF should not need OCR"))
    monkeypatch.setattr(parser, "_ocr_page", ocr)
    assert "Diffusion moves particles" in parser.extract_pdf_pages(output.getvalue())[0]
    ocr.assert_not_called()


def test_real_scanned_pdf_ocr(settings):
    from PIL import Image, ImageDraw, ImageFont

    # Generated locally; no external fixture downloads or LLM calls.
    with Image.new("RGB", (1300, 240), "white") as image:
        draw = ImageDraw.Draw(image)
        font = ImageFont.load_default(size=40)
        draw.text((30, 60), "DIFFUSION MOVES PARTICLES", fill="black", font=font)
        draw.text((30, 120), "DOWN A CONCENTRATION GRADIENT", fill="black", font=font)
        output = BytesIO()
        image.save(output, format="PDF", resolution=100)
    text = parser.extract_pdf_pages(output.getvalue())[0].upper()
    assert "DIFFUSION" in text and "GRADIENT" in text


@pytest.mark.parametrize(("text", "expected"), [
    ("Diffusion moves particles from high to low concentration.", "en"),
    ("هذا المستند يشرح كيف تنتقل الجسيمات من تركيز مرتفع إلى تركيز منخفض.", "ar"),
    ("Hii ni somo ambalo linaeleza kwa nini maji ni muhimu katika maisha ya kila siku.", "sw"),
])
def test_supported_language_detection(text, expected):
    assert document_parser.detect_language(text) == expected


def test_real_docx_extracts_paragraphs_and_tables():
    from docx import Document

    document = Document()
    document.add_paragraph("Usafirishaji wa chembe katika seli")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Aina"
    table.rows[0].cells[1].text = "Maelezo"
    output = BytesIO()
    document.save(output)

    pages, language = document_parser.extract_document_pages(output.getvalue(), "biology.docx")

    assert pages == ["Usafirishaji wa chembe katika seli\nAina | Maelezo"]
    assert language == "en"


@pytest.mark.parametrize(("filename", "content", "message"), [
    ("notes.txt", b"notes", "PDF or DOCX"),
    ("notes.docx", b"not a zip", "signature"),
])
def test_document_parser_rejects_unsupported_or_invalid_files(filename, content, message):
    with pytest.raises(HTTPException, match=message):
        document_parser.extract_document_pages(content, filename)